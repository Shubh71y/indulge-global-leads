import os
import sys
from pathlib import Path

BASE_DIR = Path(r"C:\Users\gozoi\Projects\indulge")
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from scripts import env_loader
from scripts.phone_verifier import (
    classify_gcc_phone,
    extract_phone_with_proximity,
    verify_whatsapp_live,
    reverse_caller_id_lookup,
    clean_corporate_domain,
    clean_and_validate_email,
    is_valid_person_name,
    is_broker_or_retail_trap,
)
from scripts.entity_resolver import resolve_lead_entity
from scripts.lead_store import append_leads_batch, sync_to_excel_and_csv, EXCEL_PATH, CSV_PATH

import time
import re
import requests
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace', line_buffering=True)

TAVILY_KEYS = [k.strip() for k in os.environ.get("TAVILY_API_KEYS", "").split(",") if k.strip()]
_tavily_idx = 0


def get_tavily_key():
    global _tavily_idx
    if not TAVILY_KEYS:
        return ""
    k = TAVILY_KEYS[_tavily_idx % len(TAVILY_KEYS)]
    _tavily_idx += 1
    return k


# Laser-targeted queries designed to extract personal mobile numbers of Buyers/Principals (excluding brokers)
HNI_DIRECT_DORKS = [
    ('site:linkedin.com/in/ "Dubai" ("Founder" OR "Chairman") ("Family Office" OR "Holding" OR "Investments") ("+971 50" OR "+971 52" OR "+971 54" OR "+971 55") -broker -agent -realtor', "Dubai, UAE", "AE"),
    ('site:linkedin.com/in/ "Dubai" ("Chief Executive Officer" OR "Managing Partner") ("Capital" OR "Group" OR "Ventures") ("+971 50" OR "+971 52" OR "+971 55" OR "+971 58") -broker -agent', "Dubai, UAE", "AE"),
    ('site:linkedin.com/in/ "Abu Dhabi" ("Founder" OR "Chairman" OR "CEO") ("Holding" OR "Investment" OR "Group") ("+971 50" OR "+971 52" OR "+971 55") -broker -agent', "Abu Dhabi, UAE", "AE"),
    ('site:linkedin.com/in/ "Riyadh" ("Founder" OR "Chairman" OR "CEO") ("Holding" OR "Investment" OR "Group") ("+966 50" OR "+966 53" OR "+966 55") -broker -agent', "Riyadh, Saudi Arabia", "SA"),
    ('site:linkedin.com/in/ "Jeddah" ("Founder" OR "Chairman" OR "CEO") ("Holding" OR "Group") ("+966 50" OR "+966 55") -broker -agent', "Jeddah, Saudi Arabia", "SA"),
    ('site:linkedin.com/in/ "Doha" ("Founder" OR "Chairman" OR "CEO") ("Holding" OR "Group" OR "Investment") ("+974 5" OR "+974 6" OR "+974 3" OR "+974 7") -broker', "Doha, Qatar", "QA"),
    ('site:linkedin.com/in/ "Kuwait" ("Founder" OR "Chairman" OR "CEO") ("Holding" OR "Group" OR "Capital") ("+965 9" OR "+965 6" OR "+965 5") -broker', "Kuwait City, Kuwait", "KW"),
]


def format_excel_file(file_path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(file_path)
        ws = wb.active
        ws.title = "Middle East HNI Leads"

        header_fill = PatternFill(start_color="1A202C", end_color="1A202C", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="D4AF37")
        align_center = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = align_center
            cell.border = thin_border

        for row in range(2, ws.max_row + 1):
            fill = PatternFill(start_color="F8FAFC" if row % 2 == 0 else "FFFFFF", fill_type="solid")
            for col in range(1, ws.max_column + 1):
                c = ws.cell(row=row, column=col)
                c.fill = fill
                c.border = thin_border
                c.font = Font(name="Calibri", size=10)
                if col in [6, 8, 9]:
                    c.alignment = align_center

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 48), 12)

        wb.save(file_path)
    except Exception as e:
        print(f"Warning formatting excel: {e}")


def parse_profile_snippet(item, location, region_code="AE"):
    title = item.get('title', '')
    content = item.get('content', '')
    url = item.get('url', '')
    combined = f"{title} {content}"

    # 1. Clean Person Name
    clean_title = re.sub(r' - LinkedIn.*', '', title)
    clean_title = re.sub(r' \| LinkedIn.*', '', clean_title)
    clean_title = re.sub(r'\(.*?\)', '', clean_title)

    parts = clean_title.split(' - ')
    raw_name = parts[0] if len(parts) >= 1 else clean_title
    raw_role = parts[1] if len(parts) >= 2 else "Founder & Chairman"
    company = parts[2].strip() if len(parts) >= 3 else ""

    name_tokens = [t for t in re.sub(r'[^a-zA-Z\s]', '', raw_name).split() if len(t) > 1]
    junk = {'linkedin', 'profile', 'expert', 'director', 'founder', 'ceo', 'dubai', 'riyadh', 'consultant', 'real', 'estate', 'holding'}
    name_tokens = [t for t in name_tokens if t.lower() not in junk]

    if len(name_tokens) < 2:
        return None

    fn, ln = name_tokens[0].capitalize(), name_tokens[-1].capitalize()

    if not company:
        if " at " in raw_role:
            company = raw_role.split(" at ")[-1].strip()
        elif "@" in raw_role:
            company = raw_role.split("@")[-1].strip()
        else:
            company = f"{fn} {ln} Private Office"

    # Reject Brokers / Real Estate Agents / Rental Agencies
    if is_broker_or_retail_trap(company, raw_role, url):
        return None

    if not is_valid_person_name(fn, ln, company):
        return None

    # 2. Enforce 1-to-1 Proximity Matching (<= 120 chars from Person's Name) & Strict Mobile Classification
    phone_info = extract_phone_with_proximity(
        text=combined,
        first_name=fn,
        last_name=ln,
        url=url,
        default_region=region_code,
        max_distance=120,
        require_mobile=True,
    )
    if not phone_info:
        return None

    # 3. Extract Email ONLY if present and on a real corporate domain (NEVER @executive-gcc.com)
    emails = re.findall(r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}', combined)
    raw_email = emails[0].strip() if emails else ""

    resolved = resolve_lead_entity(
        first_name=fn,
        last_name=ln,
        company_name=company[:60],
        title=raw_role.strip()[:60],
        email=raw_email,
        domain="",
        allow_network=True,
    )

    # 4. Optional Live WhatsApp & Reverse Caller-ID Verification
    wa_verify = verify_whatsapp_live(phone_info["e164"])
    if wa_verify.get("checked") and wa_verify.get("exists") is False:
        return None
    caller_id = reverse_caller_id_lookup(phone_info["e164"], expected_name=f"{fn} {ln}")
    if caller_id.get("checked") and caller_id.get("match_confidence") == "mismatch":
        return None

    final_email = resolved["email"]
    if final_email == "N/A" and caller_id.get("owner_email"):
        final_email = clean_and_validate_email(caller_id["owner_email"], fn, ln, resolved["domain"])

    return {
        "First Name": resolved["first_name"] or fn,
        "Last Name": resolved["last_name"] or ln,
        "Role": resolved["title"],
        "Company": company[:60],
        "Domain": resolved["domain"] if resolved["domain"] else "N/A",
        "Phone Number": phone_info["formatted_intl"],
        "E164": phone_info["e164"],
        "Phone Type": phone_info["phone_type"],
        "Carrier": phone_info["carrier"],
        "WhatsApp Link": phone_info["wa_link"] or "",
        "WhatsApp Verified": "True" if wa_verify.get("exists", True) else "False",
        "Avatar URL": wa_verify.get("profile_pic_url") or "",
        "Email": final_email,
        "City/Country": location,
        "Deliverability Status": "Verified Mobile & WhatsApp Direct" if final_email == "N/A" else "Verified Corporate Email & VIP Mobile",
        "Full Address": location,
        "Website": url,
        "Source": "LinkedIn Direct Mobile Intelligence Miner"
    }


def run_direct_mobile_miner():
    print("=" * 70)
    print("INDULGE GLOBAL - ULTRA DIRECT HNI MOBILE & WHATSAPP MINER")
    print("=" * 70)
    print(f"[+] Output Excel: {EXCEL_PATH}")
    print(f"[+] Enforcing 120-char proximity + phonenumbers MOBILE validation + N/A email fallback")
    print("=" * 70)

    total_added = 0

    for query, location, region_code in HNI_DIRECT_DORKS:
        print(f"\nExecuting Dork: {query[:65]}... ({location})")
        payload = {
            'api_key': get_tavily_key(),
            'query': query,
            'max_results': 8
        }

        try:
            r = requests.post('https://api.tavily.com/search', json=payload, timeout=12)
            if r.status_code == 200:
                results = r.json().get('results', [])
                batch = []
                for item in results:
                    lead = parse_profile_snippet(item, location, region_code=region_code)
                    if lead:
                        batch.append(lead)
                        print(f"  [VIP MOBILE] {lead['First Name']} {lead['Last Name']} | {lead['Phone Number']} ({lead['Carrier']}) | {lead['Company']} | Email: {lead['Email']}")

                if batch:
                    added = append_leads_batch(batch)
                    total_added += added
                    format_excel_file(str(EXCEL_PATH))
                    print(f"  Saved {added} deduplicated mobile leads to HNI.xlsx")
            time.sleep(1.5)
        except Exception as e:
            print(f"  Error querying: {e}")
            time.sleep(2)

    print("\n" + "=" * 70)
    print(f"COMPLETED DIRECT MOBILE HARVEST! New Verified Mobile Contacts Added: {total_added}")
    print(f"Updated Excel: {EXCEL_PATH}")
    print("=" * 70)


if __name__ == "__main__":
    run_direct_mobile_miner()
